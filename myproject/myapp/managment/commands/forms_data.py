from django.core.management.base import BaseCommand
import openpyxl # type: ignore
from myapp.models import Tovar, Zakaz, User, PunktVidachi

class Command(BaseCommand):
  help = 'Импорт данных из excel'

  def handle(self, *args, **options):
    wbT = openpyxl.load_workbook('import/Tovar.xlsx')
    ws = wbT.active
    for row in ws.iter_rows(min_row=2, values_only=True):
      if row[0]:
        Tovar.objects.get_or_create(articul = row[0])

    wbP = openpyxl.load_workbook('import/Пункты выдачи_import.xlsx')
    for row in wbP.active.iter_rows(min_row=2, values_only=True):
      if row[0]:
       PunktVidachi.objects.get_or_create(id = row[4])

    wbZ = openpyxl.load_workbook('import/Заказ_import.xlsx')
    for row in wbZ.active.iter_rows(min_row=2, values_only=True):
      if row[0]:
       Zakaz.objects.get_or_create(id=row[0])